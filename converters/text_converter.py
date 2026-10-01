import os
import re
import json
import yaml
import html
from pathlib import Path
from html.parser import HTMLParser
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import qrcode
from gtts import gTTS
import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Font configuration
FONT_NAME = "Helvetica"
try:
    font_paths = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    ]
    for fp in font_paths:
        if os.path.exists(fp):
            pdfmetrics.registerFont(TTFont("CustomTextFont", fp))
            FONT_NAME = "CustomTextFont"
            break
except Exception:
    pass

def parse_html_document(html_content: str):
    """
    Intelligently extracts all content from HTML, including:
    1. Static HTML elements (headings, paragraphs, lists, tables)
    2. Embedded JSON data islands (e.g. restaurant menus, catalogs, products)
    """
    json_data = None
    json_scripts = re.findall(r'<script[^>]*type=["\']application/json["\'][^>]*>(.*?)</script>', html_content, re.DOTALL | re.IGNORECASE)
    for js in json_scripts:
        try:
            parsed = json.loads(js.strip())
            if isinstance(parsed, dict) and any(k in parsed for k in ["categories", "dishes", "menu", "items", "products", "restaurant"]):
                json_data = parsed
                break
        except Exception:
            pass

    title = ""
    m_title = re.search(r'<title[^>]*>(.*?)</title>', html_content, re.IGNORECASE | re.DOTALL)
    if m_title:
        title = html.unescape(m_title.group(1).strip())
    if not title and json_data and isinstance(json_data, dict):
        rest = json_data.get("restaurant", {})
        if isinstance(rest, dict):
            title = f"{rest.get('name', '')}{rest.get('nameAccent', '')}".strip()

    class StructuredHTMLParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.ignore = False
            self.elements = []
            self.current_text = []
            self.in_table = False
            self.table_data = []
            self.current_row = []
            self.current_cell = []

        def handle_starttag(self, tag, attrs):
            tag = tag.lower()
            if tag in ('script', 'style', 'head', 'noscript', 'svg', 'meta', 'link'):
                self.ignore = True
                return
            if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'li', 'blockquote'):
                self._flush()
            elif tag == 'table':
                self._flush()
                self.in_table = True
                self.table_data = []
            elif tag == 'tr':
                self.current_row = []
            elif tag in ('td', 'th'):
                self.current_cell = []

        def handle_endtag(self, tag):
            tag = tag.lower()
            if tag in ('script', 'style', 'head', 'noscript', 'svg', 'meta', 'link'):
                self.ignore = False
                return
            if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
                text = ' '.join(''.join(self.current_text).split())
                if text:
                    self.elements.append(('heading', int(tag[1]), text))
                self.current_text = []
            elif tag in ('p', 'blockquote'):
                text = ' '.join(''.join(self.current_text).split())
                if text:
                    self.elements.append(('p', text))
                self.current_text = []
            elif tag == 'li':
                text = ' '.join(''.join(self.current_text).split())
                if text:
                    self.elements.append(('list', text))
                self.current_text = []
            elif tag in ('td', 'th'):
                cell_text = ' '.join(''.join(self.current_cell).split())
                self.current_row.append(cell_text)
                self.current_cell = []
            elif tag == 'tr':
                if self.current_row:
                    self.table_data.append(self.current_row)
                self.current_row = []
            elif tag == 'table':
                if self.table_data:
                    self.elements.append(('table', self.table_data))
                self.in_table = False
                self.table_data = []

        def handle_data(self, data):
            if self.ignore:
                return
            if self.in_table:
                self.current_cell.append(data)
            else:
                self.current_text.append(data)

        def _flush(self):
            text = ' '.join(''.join(self.current_text).split())
            if text and len(text) > 1:
                self.elements.append(('p', text))
            self.current_text = []

    parser = StructuredHTMLParser()
    parser.feed(html_content)
    parser._flush()

    return {
        "title": title or "Hujjat",
        "json_data": json_data,
        "dom_elements": parser.elements
    }

def html_to_clean_text(html_content: str) -> str:
    """Extracts clean plain text from HTML, preserving menu data and content"""
    parsed = parse_html_document(html_content)
    lines = []
    if parsed["title"]:
        lines.append(f"=== {parsed['title']} ===")
    
    if parsed["json_data"] and isinstance(parsed["json_data"], dict):
        jd = parsed["json_data"]
        if "ribbon" in jd and isinstance(jd["ribbon"], list):
            lines.append(" • ".join(jd["ribbon"]))
            lines.append("")
        for cat in jd.get("categories", []):
            lines.append(f"\n--- {cat.get('name', '')} ---")
            for d in cat.get("dishes", []):
                p = d.get("price", "")
                p_str = f"{p:,} so'm".replace(",", " ") if isinstance(p, (int, float)) else str(p)
                desc = f" ({d.get('desc')})" if d.get('desc') else ""
                lines.append(f"• {d.get('name')}: {p_str}{desc}")
                
    if parsed["dom_elements"]:
        lines.append("\n=== Ma'lumotlar ===")
        for elem_type, *args in parsed["dom_elements"]:
            if elem_type == "heading":
                lines.append(f"\n[{args[1]}] {args[0]}")
            elif elem_type == "p":
                if args[0].lower() != parsed["title"].lower():
                    lines.append(args[0])
            elif elem_type == "list":
                lines.append(f"• {args[0]}")
            elif elem_type == "table":
                for row in args[0]:
                    lines.append(" | ".join(row))
                    
    return "\n".join(lines).strip()

def html_to_docx(html_content: str, output_path: str, title: str = "") -> str:
    """Converts HTML page and embedded data to high quality Word document"""
    parsed = parse_html_document(html_content)
    doc_title = title or parsed["title"] or "Hujjat"
    
    doc = docx.Document()
    doc.add_heading(doc_title, level=0)
    
    # 1. Embedded JSON data (Dishes / Menu / Products)
    if parsed["json_data"] and isinstance(parsed["json_data"], dict):
        jd = parsed["json_data"]
        if "ribbon" in jd and isinstance(jd["ribbon"], list):
            doc.add_paragraph(" • ".join(jd["ribbon"]))
            
        cats = jd.get("categories", [])
        for cat in cats:
            cname = cat.get("name", "")
            doc.add_heading(cname, level=1)
            dishes = cat.get("dishes", [])
            if dishes:
                tbl = doc.add_table(rows=1, cols=3)
                tbl.style = 'Light Shading Accent 1'
                hdr_cells = tbl.rows[0].cells
                hdr_cells[0].text = "Taom / Mahsulot"
                hdr_cells[1].text = "Tavsifi"
                hdr_cells[2].text = "Narxi"
                
                for d in dishes:
                    row_cells = tbl.add_row().cells
                    row_cells[0].text = str(d.get("name", ""))
                    row_cells[1].text = str(d.get("desc", ""))
                    price = d.get("price", "")
                    row_cells[2].text = f"{price:,} so'm".replace(",", " ") if isinstance(price, (int, float)) else str(price)
            doc.add_paragraph("")
            
    # 2. DOM Elements (headings, paragraphs, tables, info)
    for elem_type, *args in parsed["dom_elements"]:
        if elem_type == "heading":
            level, text = args
            doc.add_heading(text, level=min(level, 3))
        elif elem_type == "p":
            text = args[0]
            if text.lower() != doc_title.lower():
                doc.add_paragraph(text)
        elif elem_type == "list":
            doc.add_paragraph(args[0], style='List Bullet')
        elif elem_type == "table":
            table_data = args[0]
            if table_data:
                rows = len(table_data)
                cols = max(len(r) for r in table_data) if rows else 0
                if rows > 0 and cols > 0:
                    tbl = doc.add_table(rows=rows, cols=cols)
                    tbl.style = 'Table Grid'
                    for r_idx, row in enumerate(table_data):
                        for c_idx, val in enumerate(row):
                            if c_idx < cols:
                                tbl.cell(r_idx, c_idx).text = val

    doc.save(output_path)
    return output_path

def html_to_pdf(html_content: str, output_path: str, title: str = "") -> str:
    """Converts HTML and embedded data to high quality styled A4 PDF"""
    parsed = parse_html_document(html_content)
    doc_title = title or parsed["title"] or "Hujjat"
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'HtmlTitle',
        parent=styles['Title'],
        fontName=FONT_NAME,
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        alignment=0,
        spaceAfter=12
    )
    h1_style = ParagraphStyle(
        'HtmlH1',
        parent=styles['Heading1'],
        fontName=FONT_NAME,
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#0f766e'),
        spaceBefore=12,
        spaceAfter=6
    )
    normal_style = ParagraphStyle(
        'HtmlNormal',
        parent=styles['Normal'],
        fontName=FONT_NAME,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1f2937')
    )
    th_style = ParagraphStyle(
        'HtmlTH',
        parent=styles['Normal'],
        fontName=FONT_NAME,
        fontSize=10,
        leading=12,
        textColor=colors.white,
    )
    
    story = [Paragraph(html.escape(doc_title), title_style)]
    
    if parsed["json_data"] and isinstance(parsed["json_data"], dict):
        jd = parsed["json_data"]
        if "ribbon" in jd and isinstance(jd["ribbon"], list):
            story.append(Paragraph(html.escape(" • ".join(jd["ribbon"])), normal_style))
            story.append(Spacer(1, 8))
            
        for cat in jd.get("categories", []):
            cname = cat.get("name", "")
            story.append(Paragraph(html.escape(cname), h1_style))
            
            dishes = cat.get("dishes", [])
            if dishes:
                table_data = [[
                    Paragraph("<b>Taom / Mahsulot</b>", th_style),
                    Paragraph("<b>Tavsifi</b>", th_style),
                    Paragraph("<b>Narxi</b>", th_style)
                ]]
                for d in dishes:
                    price = d.get("price", "")
                    p_str = f"{price:,} so'm".replace(",", " ") if isinstance(price, (int, float)) else str(price)
                    table_data.append([
                        Paragraph(html.escape(str(d.get("name", ""))), normal_style),
                        Paragraph(html.escape(str(d.get("desc", ""))), normal_style),
                        Paragraph(html.escape(p_str), normal_style)
                    ])
                
                t = Table(table_data, colWidths=[140, 260, 115])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
                    ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
                    ('TOPPADDING', (0,0), (-1,-1), 5),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))

    if parsed["dom_elements"]:
        story.append(Paragraph("<b>Ma'lumotlar va aloqa:</b>", h1_style))
        for elem_type, *args in parsed["dom_elements"]:
            if elem_type == "heading":
                story.append(Paragraph(f"<b>{html.escape(args[1])}</b>", h1_style))
            elif elem_type == "p":
                if args[0].lower() != doc_title.lower():
                    story.append(Paragraph(html.escape(args[0]), normal_style))
                    story.append(Spacer(1, 3))
            elif elem_type == "list":
                story.append(Paragraph(f"• {html.escape(args[0])}", normal_style))
                story.append(Spacer(1, 2))

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    doc.build(story)
    return output_path

def text_to_pdf(text: str, output_path: str, title: str = "Hujjat") -> str:
    """Converts raw text or HTML to styled A4 PDF"""
    if "<html" in text[:500].lower() or "<!doctype html" in text[:500].lower():
        return html_to_pdf(text, output_path, title=title)
        
    styles = getSampleStyleSheet()
    normal_style = ParagraphStyle(
        'TxtNormal',
        parent=styles['Normal'],
        fontName=FONT_NAME,
        fontSize=11,
        leading=16,
        textColor=colors.HexColor('#1f2937')
    )
    title_style = ParagraphStyle(
        'TxtTitle',
        parent=styles['Title'],
        fontName=FONT_NAME,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1e40af'),
        spaceAfter=15
    )
    
    pdf = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    story = [Paragraph(html.escape(title), title_style)]
    for line in text.split("\n"):
        line = line.strip()
        if line:
            story.append(Paragraph(html.escape(line), normal_style))
            story.append(Spacer(1, 6))
        else:
            story.append(Spacer(1, 8))
            
    pdf.build(story)
    return output_path

def text_to_docx(text: str, output_path: str, title: str = "Hujjat") -> str:
    """Converts raw text or HTML to Word (.docx) document"""
    if "<html" in text[:500].lower() or "<!doctype html" in text[:500].lower():
        return html_to_docx(text, output_path, title=title)

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(output_path)
    return output_path

def text_to_speech(text: str, output_path: str, lang: str = "tr") -> str:
    """Converts text to speech MP3 audio"""
    if "<html" in text[:500].lower() or "<!doctype html" in text[:500].lower():
        text = html_to_clean_text(text)
    tts = gTTS(text=text[:3000], lang=lang, slow=False)
    tts.save(output_path)
    return output_path

def text_to_qr(text: str, output_path: str) -> str:
    """Generates high quality QR code image from text or URL"""
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=3,
    )
    qr.add_data(text)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#1e293b", back_color="white")
    img.save(output_path)
    return output_path

def json_to_csv(input_path: str, output_path: str) -> str:
    """Converts JSON array/object to CSV"""
    with open(input_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        df = pd.DataFrame(data)
    else:
        df = pd.json_normalize(data)
    df.to_csv(output_path, index=False, encoding="utf-8-sig")
    return output_path

def json_to_yaml(input_path: str, output_path: str) -> str:
    """Converts JSON to YAML"""
    with open(input_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    with open(output_path, "w", encoding="utf-8") as f:
        yaml.dump(data, f, allow_unicode=True, sort_keys=False)
    return output_path

def yaml_to_json(input_path: str, output_path: str) -> str:
    """Converts YAML to JSON"""
    with open(input_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return output_path
